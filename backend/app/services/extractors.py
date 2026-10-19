"""Every kind of source → the page/block shape the PDF parser produces.

    pages = [{"page_number": int, "blocks": [{"kind": "heading"|"body"|"code"|"table", "text": str}]}]

Chunking, embedding, search and citations then work the same for all of them.
PowerPoint slides are pages; Word, Markdown and text sources are a single page and
are cited by section title instead. Tables become Markdown tables.

`validate` runs in the upload request and rejects files that aren't what their
extension says, Office files that would expand to too much data (zip bombs), and
macro-enabled documents.
"""
import io
import re
import zipfile
from collections.abc import Iterable, Iterator

from app.config import settings
from app.services.content_security import UnsafeSourceError, validate_pdf
from app.services.pdf_parser import extract_pages_from_pdf

KINDS = ("pdf", "docx", "pptx", "txt", "md", "text")  # "text" = pasted text
EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".txt": "txt", ".md": "md", ".markdown": "md"}
LABELS = {"pdf": "PDF", "docx": "Word document", "pptx": "PowerPoint file", "txt": "text file", "md": "Markdown file", "text": "text"}
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "txt": "text/plain; charset=utf-8",
    "md": "text/markdown; charset=utf-8",
    "text": "text/plain; charset=utf-8",
}
SUPPORTED_MESSAGE = "Supported files: PDF, Word (.docx), PowerPoint (.pptx), Markdown (.md) and text (.txt)."

MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
MAX_ZIP_ENTRIES = 5000
MAX_SLIDES = 500
MAX_TEXT_CHARS = 2_000_000
_MAIN_PART = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml"}
_SLIDE_PART = re.compile(r"ppt/slides/slide\d+\.xml")


def kind_for_filename(filename: str) -> str | None:
    """The source kind for an uploaded file's name, or None when it isn't supported."""
    for ext, kind in EXTENSIONS.items():
        if filename.lower().endswith(ext):
            return kind
    return None


# --- Validation -----------------------------------------------------------------

def _validate_office(kind: str, data: bytes) -> None:
    label = LABELS[kind]
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
    except (zipfile.BadZipFile, ValueError) as exc:
        raise UnsafeSourceError(f"That file isn't a valid {label}.") from exc
    names = {entry.filename for entry in entries}
    if _MAIN_PART[kind] not in names:
        raise UnsafeSourceError(f"That file isn't a valid {label}.")
    # Declared sizes bound what zipfile will ever decompress, so checking them stops zip bombs.
    if len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES:
        raise UnsafeSourceError(f"That {label} expands to too much data to process safely.")
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise UnsafeSourceError(f"{label[0].upper()}{label[1:]}s containing macros aren't accepted for safety.")
    if kind == "pptx" and sum(1 for name in names if _SLIDE_PART.fullmatch(name)) > MAX_SLIDES:
        raise UnsafeSourceError(f"PowerPoint files may contain at most {MAX_SLIDES} slides.")


def decode_text(data: bytes) -> str:
    """UTF-8 (with or without BOM), else Windows-1252. Raises UnsafeSourceError for binary data."""
    if b"\x00" in data:
        raise UnsafeSourceError("That file isn't plain text.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def validate(kind: str, data: bytes) -> None:
    """Reject a file that isn't a safe, readable source of this kind (user-facing messages)."""
    if kind == "pdf":
        if not data.startswith(b"%PDF-"):
            raise UnsafeSourceError("That file isn't a valid PDF")
        validate_pdf(data, max_pages=settings.MAX_PDF_PAGES)
    elif kind in _MAIN_PART:
        _validate_office(kind, data)
    else:
        text = decode_text(data)
        if not text.strip():
            raise UnsafeSourceError("That file is empty.")
        if len(text) > MAX_TEXT_CHARS:
            raise UnsafeSourceError(f"Text sources may contain at most {MAX_TEXT_CHARS:,} characters.")


# --- Extraction --------------------------------------------------------------------

def _cell(text: str) -> str:
    return " ".join(text.split()).replace("|", "\\|")


def markdown_table(rows: list[list[str]]) -> str:
    """Rows of cell text → a Markdown table; the first row is the header."""
    rows = [[_cell(c) for c in row] for row in rows if any(c.strip() for c in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    lines = ["|" + "|".join(rows[0]) + "|", "|" + "|".join(["---"] * width) + "|"]
    lines += ["|" + "|".join(row) + "|" for row in rows[1:]]
    return "\n".join(lines)


def _docx_pages(data: bytes) -> list[dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    blocks = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            try:
                style = (paragraph.style.name if paragraph.style is not None else "") or ""
            except (KeyError, AttributeError, ValueError):  # a style the document references but doesn't define
                style = ""
            heading = style == "Title" or style.startswith("Heading")
            blocks.append({"kind": "heading" if heading else "body", "text": text})
        elif tag == "tbl":
            table = Table(child, document)
            markdown = markdown_table([[cell.text for cell in row.cells] for row in table.rows])
            if markdown:
                blocks.append({"kind": "table", "text": markdown})
    return [{"page_number": 1, "blocks": blocks}] if blocks else []


def _pptx_shapes(shapes) -> Iterator:
    """Shapes in reading order (top to bottom, left to right), with groups flattened."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def position(shape) -> tuple[int, int]:
        return (shape.top or 0, shape.left or 0)

    for shape in sorted(shapes, key=position):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shapes(shape.shapes)
        else:
            yield shape


def _frame_text(frame) -> str:
    return "\n".join(p.text.strip() for p in frame.paragraphs if p.text.strip())


def _pptx_pages(data: bytes) -> tuple[int, list[dict]]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    pages = []
    slides = list(presentation.slides)
    for number, slide in enumerate(slides, start=1):
        title_shape = slide.shapes.title
        blocks = []
        title = _frame_text(title_shape.text_frame) if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            blocks.append({"kind": "heading", "text": " ".join(title.split())})
        for shape in _pptx_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if markdown := markdown_table(rows):
                    blocks.append({"kind": "table", "text": markdown})
            elif shape.has_text_frame and (text := _frame_text(shape.text_frame)):
                blocks.append({"kind": "body", "text": text})
        if slide.has_notes_slide and (notes := _frame_text(slide.notes_slide.notes_text_frame)):
            blocks.append({"kind": "body", "text": f"Notes: {notes}"})
        if blocks:
            pages.append({"page_number": number, "blocks": blocks})
    return len(slides), pages


def _paragraphs(lines: Iterable[str]) -> Iterator[str]:
    """Blank-line separated paragraphs."""
    current: list[str] = []
    for line in lines:
        if line.strip():
            current.append(line.rstrip())
        elif current:
            yield "\n".join(current)
            current = []
    if current:
        yield "\n".join(current)


_MD_HEADING = re.compile(r"^#{1,6}\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def _markdown_blocks(text: str) -> list[dict]:
    """Headings, fenced code, pipe tables and paragraphs, in order. Inline syntax is kept."""
    blocks: list[dict] = []
    pending: list[str] = []  # prose lines waiting to become paragraphs

    def flush() -> None:
        blocks.extend({"kind": "body", "text": p} for p in _paragraphs(pending))
        pending.clear()

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if heading := _MD_HEADING.match(line):
            flush()
            blocks.append({"kind": "heading", "text": heading.group(1)})
            i += 1
        elif fence := _MD_FENCE.match(line):
            flush()
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith(fence.group(1)):
                j += 1
            code = "\n".join(lines[i + 1:j]).strip("\n")
            if code.strip():
                blocks.append({"kind": "code", "text": code})
            i = j + 1
        elif line.lstrip().startswith("|"):
            flush()
            j = i
            while j < len(lines) and lines[j].lstrip().startswith("|"):
                j += 1
            blocks.append({"kind": "table", "text": "\n".join(row.strip() for row in lines[i:j])})
            i = j
        else:
            pending.append(line)
            i += 1
    flush()
    return blocks


def extract(kind: str, data: bytes) -> tuple[int, list[dict], int]:
    """Source bytes → (page_count, pages, ocr_page_count). Pages carry blocks; their text
    is assembled after sanitizing (see ingest)."""
    if kind == "pdf":
        page_count, pages, ocr_pages = extract_pages_from_pdf(data)
        return page_count, [{"page_number": p["page_number"], "blocks": p["blocks"]} for p in pages], ocr_pages
    if kind == "docx":
        return 1, _docx_pages(data), 0
    if kind == "pptx":
        slide_count, pages = _pptx_pages(data)
        return slide_count, pages, 0
    text = decode_text(data)
    if kind == "md":
        blocks = _markdown_blocks(text)
    else:
        blocks = [{"kind": "body", "text": p} for p in _paragraphs(text.split("\n"))]
    return 1, ([{"page_number": 1, "blocks": blocks}] if blocks else []), 0
